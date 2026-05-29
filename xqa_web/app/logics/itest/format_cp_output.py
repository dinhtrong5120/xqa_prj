import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill


def is_all_empty(row, columns):
    return all(row.iloc[col] == "" for col in columns)


def condense_data(df, column_trai, column_phai, column_giua):
    columns_trai_phai = column_trai + column_phai

    for index in range(0, len(df), 1):
        if index < 100:
            while is_all_empty(df.iloc[index], columns_trai_phai):
                j = 0
                for x in range(1, index):
                    # print('x: ', x)
                    if (df.iloc[index - x, column_giua] == '').all() and (df.iloc[index - x, column_trai] != '').any():
                        j = x
                        # print('index - x -1: ', index - x - 1)
                    elif (df.iloc[index - x, column_giua] == '').all() and (
                            df.iloc[index - x, column_trai] == '').all():
                        if (df.iloc[index - x - 1, column_giua] == '').all() and (
                                df.iloc[index - x - 1, column_trai] != '').any():
                            j = x
                    elif (df.iloc[index - x, column_giua] != '').any():
                        break
                    else:
                        if j != 0:
                            j = j - 1
                        break
                if index != 1:
                    j = j - 1
                # Kiểm tra hàng phía trên
                if all(df.iat[index - j - 1, col] == "" for col in column_giua) and any(
                        df.iat[index - j - 1, col] != "" for col in column_trai):
                    # Dồn dữ liệu từ hàng hiện tại lên hàng phía trên
                    for col in column_giua:
                        df.iat[index - j - 1, col] = df.iat[index, col] if df.iat[index, col] != "" else df.iat[
                            index - j - 1, col]
                    # Xóa dữ liệu của hàng hiện tại
                    for col in column_giua:
                        df.iat[index, col] = ""
                    index -= 1  # Tiếp tục kiểm tra hàng phía trên
                else:
                    break

    # Xóa các hàng rỗng
    df.dropna(how='all', inplace=True)

    return df.reset_index(drop=True)


def process_compare_output(file_path):
    wb = Workbook()
    # file_name = 'condensed_data.xlsx'
    all_sheets = pd.ExcelFile(file_path).sheet_names
    for item in all_sheets:
        df = pd.read_excel(file_path, sheet_name=item)
        df = df.fillna("")
        strart_colum_1_bante = df.columns.get_loc('１番手')
        strart_colum_2_bante_mae = df.columns.get_loc('２番手削除前')
        strart_colum_2_bante_go = df.columns.get_loc('２番手削除後')
        column_trai = [i for i in range(strart_colum_1_bante, strart_colum_2_bante_mae - 2)]
        column_phai = [i for i in range(strart_colum_2_bante_go, df.shape[1])]
        column_giua = [i for i in range(strart_colum_2_bante_mae, strart_colum_2_bante_go)]
        df = condense_data(df, column_trai, column_phai, column_giua)

        df = df[~(df == '').all(axis=1)]
        df.reset_index(drop=True, inplace=True)
        ws = wb.create_sheet(title=item)  # Tạo sheet mới với tên như Sheet đang xét.
        # Ghi tiêu đề
        list_column = df.columns.tolist()
        list_title_excel = ['' if str(item).startswith('Unnamed') else item for item in list_column]
        ws.append(list_title_excel)
        # Đọc lại dữ liệu từ DataFrame và điền vào file Excel
        for row_index, row in df.iterrows():
            for col_index, value in enumerate(row):
                cell = ws.cell(row=row_index + 2, column=col_index + 1)
                cell.value = value
        # Tạo màu tô
        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        blue_fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
        # Đọc lại và tô màu dựa trên điều kiện
        for row_index, row in df.iterrows():
            # Kiểm tra cột 3 (column_phai) nếu tất cả các ô rỗng
            if all(row.iloc[col] == "" for col in column_phai):
                # Tô màu vàng cho tất cả các ô trong cột column_phai
                for col in column_phai:
                    ws.cell(row=row_index + 2, column=col + 1).fill = yellow_fill
                # Tô màu xanh cho tất cả các ô còn lại trong column_trai và column_giua nếu rỗng
                if all(row.iloc[col] == "" for col in column_trai):
                    for col in column_trai:
                        ws.cell(row=row_index + 2, column=col + 1).fill = blue_fill

                if all(row.iloc[col] == "" for col in column_giua):
                    for col in column_giua:
                        ws.cell(row=row_index + 2, column=col + 1).fill = blue_fill
            else:
                # Kiểm tra nếu có ít nhất 1 giá trị không rỗng trong cột column_phai
                if all(row.iloc[col] == "" for col in column_trai):
                    for col in column_trai:
                        ws.cell(row=row_index + 2, column=col + 1).fill = yellow_fill

                if all(row.iloc[col] == "" for col in column_giua):
                    for col in column_giua:
                        ws.cell(row=row_index + 2, column=col + 1).fill = yellow_fill
    try:
        std = wb['Sheet']
        wb.remove(std)
    except:
        pass
    wb.save(file_path)