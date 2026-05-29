import sys
import os
import pandas as pd
import zipfile
from xqa_web.app.setting import BASE_DIR_OUTPUTS, BASE_DIR_CADICS_TEMP
from openpyxl import load_workbook
# from src.db.funtion_database import query_data


def frame_empty():
    data_empty = pd.DataFrame([[""] * 30] * 20)
    column_names = [f'{i + 1}' for i in range(30)]
    data_empty = pd.DataFrame(data_empty, columns=column_names)
    data_empty = data_empty.fillna("")
    return data_empty

def check_file_out(car, powertrain, plant, case):
    list_check = []
    car = str(car).upper()
    #working = os.path.dirname(__file__)
    folder_out = car + "_" + powertrain + "_" + plant + "_" + case
    folder_output = os.path.join(BASE_DIR_OUTPUTS, folder_out)
    folder_output = folder_output.replace("\\", "/")
    name_zip = folder_out + ".zip"
    list_file_out = ["Car配車要望表.xlsx", "WTC仕様用途一覧表.xlsx", "WTC要望集約兼チェックリスト.xlsx", "実験部品.xlsx",
                     "特性管理部品リスト.xlsx",  "車両仕様表サンプル.xlsx", "File Log.xlsx"]
    for file in list_file_out:
        link_check = folder_output + "/" + file
        if os.path.exists(link_check):
            list_check.append(link_check)
        else:
            list_check.append(None)
    return list_check, folder_output, name_zip


def read_cadics(link_cadic):
    if link_cadic is not None:
        data = pd.read_csv(link_cadic, header=None)
        data = data.fillna("")
        num_columns = len(data.columns)
        column_names = [f'{i + 1}' for i in range(num_columns)]
        data.columns = column_names
        return data, 1
    else:
        return frame_empty(), 0


def read_output(link_file, sheetname):
    print(f"link_file: {link_file}")
    print(f"sheetname: {sheetname}")
    if link_file is not None:
        wb = load_workbook(link_file)
        sheetname = sheetname if sheetname in wb.sheetnames else sheetname.split("_")[0]
        data = pd.read_excel(link_file, sheet_name=sheetname, header=None)
        data = data.fillna("")
        num_columns = len(data.columns)
        column_names = [f'{i + 1}' for i in range(num_columns)]
        data.columns = column_names
        return data
    return frame_empty()


def zip_folder(folder_path, zip_path):
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for root, _, files in os.walk(folder_path):
            for file in files:
                file_path = os.path.join(root, file)
                zip_path = os.path.relpath(file_path, folder_path)
                zipf.write(file_path, zip_path)


def write_cadic_temp1(use, pos, car, pwt, plant, case, df):
    if pos == "admin":
        #working = os.path.dirname(__file__)
        folder_name = str(use).upper() + "_" + str(car).upper() + "_" + pwt + "_" + plant + "_" + case
        folder_data = os.path.join("./cadic_temp", folder_name)
        folder_data = folder_data.replace("\\", "/")
        link_cadic = folder_data + "/CADICS_ALL.csv"
        if not os.path.exists(folder_data):
            os.mkdir(folder_data)
        df.to_csv(link_cadic, index=None, header=None)
def write_cadic_temp(use, pos, car, pwt, plant, case, df):
    if pos in ["admin", "master"]:
        #working = os.path.dirname(__file__)
        folder_name = str(use).upper() + "_" + str(car).upper() + "_" + pwt + "_" + plant + "_" + case
        folder_data = os.path.join(BASE_DIR_CADICS_TEMP, folder_name)
        folder_data = folder_data.replace("\\", "/")
        link_cadic = folder_data + "/CADICS_ALL.csv"
        if not os.path.exists(folder_data):
            os.mkdir(folder_data)
        df.to_csv(link_cadic, index=None, header=None)
