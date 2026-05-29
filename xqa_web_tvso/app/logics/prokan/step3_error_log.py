# error_log.py
import re
import pandas as pd
import numpy as np

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import os


LOG_COLUMNS = [
    "Error Type",
    "From",
    "Sheet",
    "Index Row",
    "Index Column",
    "Value Error",
    "Detail Error"
]

ERROR_TYPE_COLOR_MAP = {
    "Duplicate Spec Code":      "FFFF99",  # Vàng nhạt
    "Spec Code Not In Car":     "99CCFF",  # Xanh nhạt
    "Spec Code Not In Form":    "FFCC99",  # Cam nhạt
    "Spec Code Format Error":   "CCCCCC",  # Xám nhạt
    "Invalid Spec Code Value":  "FF9999",  # Đỏ nhạt
    "Car Sheet Format Error":   "CC99FF",  # Tím nhạt
    # Thêm các loại khác nếu cần
}

def export_log_with_color_to_sheet(df_log, file_path, sheet_name):
    # If file does not exist, create new one
    if not os.path.exists(file_path):
        with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
            df_log.to_excel(writer, index=False, sheet_name=sheet_name)
    else:
        # Try to open as Excel, if fails, recreate it
        try:
            wb = load_workbook(file_path)
        except Exception as e:
            print(f"Warning: {file_path} is not a valid Excel file. Re-creating it. Error: {e}")
            os.remove(file_path)
            with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
                df_log.to_excel(writer, index=False, sheet_name=sheet_name)
            wb = load_workbook(file_path)
        # If sheet exists, remove it
        if sheet_name in wb.sheetnames:
            std = wb[sheet_name]
            wb.remove(std)
        # Write new sheet (append mode)
        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            df_log.to_excel(writer, index=False, sheet_name=sheet_name)

    # Color rows by error type
    wb = load_workbook(file_path)
    ws = wb[sheet_name]
    for i, row in df_log.iterrows():
        err_type = row["Error Type"]
        color = ERROR_TYPE_COLOR_MAP.get(err_type, None)
        if color:
            fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            for col in range(1, len(df_log.columns)+1):
                ws.cell(row=i+2, column=col).fill = fill  # +2 for header row
    wb.save(file_path)

def collect_error_log():
    """
    Tạo DataFrame log lỗi.
    """
    return pd.DataFrame(columns=LOG_COLUMNS)

def log_duplicate_spec_code(df_form, sheet_name, spec_col=10, start_row=37):
    """
    Kiểm tra và log các spec code trùng lặp trong form.
    Trả về list dict log.
    """
    logs = []
    codes = df_form.iloc[start_row:, spec_col].dropna().astype(str).map(str.strip)
    code_counts = codes.value_counts()
    dup_codes = code_counts[code_counts > 1]
    for code, count in dup_codes.items():
        # Lấy tất cả vị trí dòng của code này
        rows = codes[codes == code].index.tolist()
        for idx, row in enumerate(rows):
            logs.append({
                "Error Type": "Duplicate Spec Code",
                "From": "Form",
                "Sheet": sheet_name,
                "Index Row": row,
                "Index Column": spec_col,
                "Value Error": code,
                "Detail Error": f"Spec code '{code}' is duplicated in Form sheet '{sheet_name}' at rows {rows}, column {spec_col}."
            })
    return logs

def log_spec_code_not_in_car(ls_form_codes, ls_car_codes, sheet_name, form_rows, spec_col=10):
    """
    Log các spec code có trong form nhưng không có trong car.
    """
    logs = []
    car_codes_set = set([str(c).strip().lower() for c in ls_car_codes])
    for i, code in enumerate(ls_form_codes):
        code_std = str(code).strip().lower()
        if code_std and code_std not in car_codes_set:
            logs.append({
                "Error Type": "Spec Code Not In Car",
                "From": "Form",
                "Sheet": sheet_name,
                "Index Row": form_rows[i],
                "Index Column": spec_col,
                "Value Error": code,
                "Detail Error": f"Spec code '{code}' in Form sheet '{sheet_name}' at row {form_rows[i]}, column {spec_col} does not exist in Car sheet."
            })
    return logs

def log_spec_code_not_in_form(ls_car_codes, ls_form_codes, sheet_name, car_rows, spec_col=136):
    """
    Log các spec code có trong car nhưng không có trong form.
    Xử lý trường hợp 1 cell car có nhiều spec code cách nhau bởi dấu ','.
    """
    logs = []
    form_codes_set = set([str(c).strip().lower() for c in ls_form_codes])
    for i, code in enumerate(ls_car_codes):
        if not isinstance(code, str):
            code = str(code)
        # Tách spec code theo dấu phẩy
        code_list = [c.strip() for c in code.split(",") if c.strip()]
        for sub_code in code_list:
            code_std = sub_code.strip().lower()
            if code_std and code_std not in form_codes_set:
                logs.append({
                    "Error Type": "Spec Code Not In Form",
                    "From": "Car",
                    "Sheet": sheet_name,
                    "Index Row": car_rows[i],
                    "Index Column": spec_col,
                    "Value Error": sub_code,
                    "Detail Error": f"Spec code '{sub_code}' in Car sheet '{sheet_name}' at row {car_rows[i]}, column {spec_col} does not exist in Form sheet."
                })
    return logs

def log_spec_code_format_error(codes, sheet_name, rows, spec_col, regex_format=r"^[A-Za-z0-9\*_]+$"):
    """
    Log các spec code có định dạng bất thường.
    - codes: list các spec code cần kiểm tra.
    - regex_format: mẫu hợp lệ.
    """
    logs = []
    for i, code in enumerate(codes):
        code_str = str(code).strip()
        if code_str and not re.match(regex_format, code_str):
            logs.append({
                "Error Type": "Spec Code Format Error",
                "From": "Form",
                "Sheet": sheet_name,
                "Index Row": rows[i],
                "Index Column": spec_col,
                "Value Error": code,
                "Detail Error": f"Spec code '{code}' in Form sheet '{sheet_name}' at row {rows[i]}, column {spec_col} has an abnormal format."
            })
    return logs

def log_invalid_spec_code_value(df, sheet_name, allowed_values=("", None, "*", "-", np.nan), from_type="Car", key_col=136):
    """
    Log các giá trị spec code không hợp lệ trong DataFrame feature matrix.
    - Kiểm tra tất cả cell (trừ cột key spec code).
    - Chỉ chấp nhận số, "", None, "*", "-", np.nan.
    """
    logs = []
    # Lấy tên cột key
    if isinstance(key_col, int):
        if key_col in df.columns:
            key_col_name = key_col
        else:
            # Nếu cột là index, có thể là tên cột
            key_col_name = df.columns[0]
    else:
        key_col_name = key_col

    # Duyệt từng cell, bỏ qua cột key spec code
    for row_idx in df.index:
        for col in df.columns:
            if col == key_col_name:
                continue  # bỏ qua cột spec code
            val = df.at[row_idx, col]
            # Nếu là số hợp lệ
            if isinstance(val, (int, float)) and not pd.isna(val):
                continue
            # Nếu là NaN
            if pd.isna(val):
                continue
            # Nếu là string thì strip và kiểm tra allowed, hoặc là số nguyên dương dạng chuỗi
            if isinstance(val, str):
                v = val.strip()
                if v in allowed_values or v.isdigit():
                    continue
            # Hợp lệ nếu là allowed_values (không phải string)
            elif val in allowed_values:
                continue
            # Nếu không thuộc bất kỳ trường hợp hợp lệ nào --> log lỗi
            logs.append({
                "Error Type": "Invalid Spec Code Value",
                "From": from_type,
                "Sheet": sheet_name,
                "Index Row": row_idx,
                "Index Column": col,
                "Value Error": val,
                "Detail Error": (
                    f"Value '{val}' for spec code in {from_type} sheet '{sheet_name}' "
                    f"at row {row_idx}, column {col} is invalid. "
                    "Expected a number, empty string, None, '*', or '-'."
                )
            })
    return logs

def log_car_sheet_format_error(errors, sheet_name):
    """
    Log các lỗi định dạng sheet car.
    """
    logs = []
    for err in errors:
        logs.append({
            "Error Type": "Car Sheet Format Error",
            "From": "Car",
            "Sheet": sheet_name,
            "Index Row": None,
            "Index Column": None,
            "Value Error": None,
            "Detail Error": f"Car sheet '{sheet_name}' format error: {err}"
        })
    return logs