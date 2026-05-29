from xqa_web_tvso.app.setting import *
from xqa_web_tvso.app.logics.itest.itest_btn_man import data_normalization
import pandas as pd
from openpyxl import load_workbook


def get_file_path_edit(folder_output_itest_path,option_output,option_file):
    path_edit = os.path.join(folder_output_itest_path,option_output)
    if os.path.isfile(path_edit):
        pass
    else:
        path_edit = os.path.join(path_edit,option_file)
    # print(path_edit)
    return path_edit


def track_changes(edited_df, original_df):
    changes_df = pd.DataFrame(columns=['Cột thay đổi', 'Hàng thay đổi', 'Địa chỉ ô', 'Giá trị cũ', 'Giá trị mới'])
    
    # for col in original_df.columns:
    for col_index in range(original_df.shape[1]):
        for idx in range(len(original_df)):
            if data_normalization(edited_df.iat[idx, col_index]) != data_normalization(original_df.iat[idx, col_index]):
                change_row = pd.DataFrame({
                    'Cột thay đổi': [col_index],
                    'Hàng thay đổi': [idx],
                    'Giá trị cũ': [original_df.iat[idx, col_index]],
                    'Giá trị mới': [edited_df.iat[idx, col_index]]
                })
                changes_df = pd.concat([changes_df, change_row], ignore_index=True)
    return changes_df


def edit_file(file_path_edit,changes_df,option_sheet):
    # print(file_path_edit)
    wb = load_workbook(file_path_edit)
    ws = wb[option_sheet]
    for index, row in changes_df.iterrows():
        col_index  = row['Cột thay đổi']
        row_index = row['Hàng thay đổi']
        new_value = row['Giá trị mới'] 
        ws.cell(row=row_index + 1, column=col_index+1, value=new_value) 
    wb.save(file_path_edit)


def save_itest(folder_output_itest_path,option_output,option_file,option_sheet, edited_df, original_df):
    # lấy đường dẫn file cần sửa
    file_path_edit = get_file_path_edit(folder_output_itest_path,option_output,option_file)
    # so sánh sự khác nhau giữa df đọc file và df đã sửa trên giao diện, trả ra hàng , cột , giá trị cũ và giá trị mới trong change_df
    changes_df = track_changes(edited_df, original_df)
    # sửa file
    # print("======================changes_df======================",changes_df)
    edit_file(file_path_edit,changes_df,option_sheet)
    message = "Save complete !!!"
    return message