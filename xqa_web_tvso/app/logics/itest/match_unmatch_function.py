import unicodedata
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill
import os
import shutil
import json
import numpy as np


def copy_folder(input_path, output_path):
    # Check if the source directory exists
    if os.path.exists(input_path):
        if os.path.exists(output_path):
            shutil.rmtree(output_path)
        # Copy the folder
        shutil.copytree(input_path, output_path)
        print(f'Folder copied from {input_path} to {output_path}')
    else:
        print(f'Source folder does not exist: {input_path}')


def normalize_japanese_text(input_text):
    normalized_text = ''
    if isinstance(input_text, str):
        for char in input_text:
            normalized_char = unicodedata.normalize('NFKC', char)
            normalized_text += normalized_char
        normalized_text = normalized_text.replace("\n", "")
        normalized_text = normalized_text.strip()
        return normalized_text
    else:
        return input_text


def fill_data(dict_part, df_file_VC, flag_same):
    dict_col = {'commodity': 1, 'partname': 2, 'bodytype': [3, 4]}
    df_file_VC_filter = df_file_VC.copy()
    for key in dict_part.keys():
        if dict_part[key] != "nan":
            if key != "bodytype":
                df_file_VC_filter = df_file_VC_filter[
                    df_file_VC_filter[dict_col[key]] == normalize_japanese_text(dict_part[key]).upper()]
            else:
                df_file_VC_filter = df_file_VC_filter.loc[df_file_VC_filter[dict_col[key]].apply(
                    lambda col: col.str.contains(normalize_japanese_text(dict_part[key]).upper(), na=False)).any(
                    axis=1)]

    for index, row in df_file_VC_filter.iterrows():
        if flag_same == True:
            df_file_VC.iloc[index] = df_file_VC.iloc[index].replace("●", 'MATCH')
            df_file_VC.iloc[index] = df_file_VC.iloc[index].replace("○", 'MATCH')
        else:
            df_file_VC.iloc[index] = df_file_VC.iloc[index].replace("●", 'UNMATCH')
            df_file_VC.iloc[index] = df_file_VC.iloc[index].replace("○", 'UNMATCH')

    return df_file_VC


def get_data_write(df_file_VC):
    address_color = []
    jikken_full_match = ''
    list_jikken_full_match = []
    for col in range(5, df_file_VC.shape[1]):
        unique_values = df_file_VC[col].unique()
        if '●' not in unique_values and "MATCH" in unique_values and '○' not in unique_values and "UNMATCH" not in unique_values:
            address_color.append(get_column_letter(col + 1) + "8")
            jikken_full_match = df_file_VC.iat[5, col]
            list_jikken_full_match.append(jikken_full_match)
        else:
            address_color.append(get_column_letter(col + 1) + "7")
    df_write = df_file_VC.loc[10:, 5:]
    return address_color, df_write, list_jikken_full_match


def write_excel(address_color, df_write, path):
    with pd.ExcelWriter(path, engine='openpyxl', mode="a", if_sheet_exists="overlay") as writer:
        fill_color = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")
        df_write.to_excel(writer, sheet_name="原紙", index=None, header=None, startcol=5, startrow=10)
        worksheet = writer.sheets['原紙']
        for address in address_color:
            worksheet[address].fill = fill_color


def change_VC_file(list_row_same, list_row_not_same, dict_part_AABB, link_folder_vc, link_folder_output):
    list_path_files_VC = [os.path.join(link_folder_vc, file) for file in os.listdir(link_folder_vc) if
                          file.endswith(".xlsx")]
    
    dict_match_jikken_trigger = dict()
    for path in list_path_files_VC:
        try:
            df_file_VC = pd.read_excel(path, sheet_name="原紙", header=None)
            df_file_VC[1] = df_file_VC[1].ffill()
            df_file_VC = df_file_VC.map(lambda x: normalize_japanese_text(x).upper() if isinstance(x, str) else x)
            # replace rei
            tmp_list = np.where(df_file_VC[0].values == "例")[0]
            if len(tmp_list) > 0:
                rei_index = tmp_list[0]   
                df_file_VC.iloc[rei_index:rei_index + 4] = df_file_VC.iloc[rei_index:rei_index + 4].replace("●", '')
                df_file_VC.iloc[rei_index:rei_index + 4] = df_file_VC.iloc[rei_index:rei_index + 4].replace("○", '')

            df_file_VC[2] = df_file_VC[2].ffill()
            df_file_VC[3] = df_file_VC[3].ffill()
            df_file_VC[4] = df_file_VC[4].ffill()
            df_file_VC = df_file_VC.fillna("").astype(str)
            
            for row_same in list_row_same:
                dict_part = dict_part_AABB[row_same]
                df_file_VC = fill_data(dict_part, df_file_VC, flag_same=True)
            address_color, df_write, list_jikken_full_match = get_data_write(df_file_VC)

            for row_not_same in list_row_not_same:
                dict_part = dict_part_AABB[row_not_same]
                df_file_VC = fill_data(dict_part, df_file_VC, flag_same=False)
            address_color, df_write, list_jikken_full_match_not_same = get_data_write(df_file_VC)
            # Write json file contain : name file have jikken full match in VC file , jikken of that file  
            file_name_full_match = os.path.basename(path)
            if len(list_jikken_full_match) > 0:
                dict_match_jikken_trigger[file_name_full_match] = list_jikken_full_match
            write_excel(address_color, df_write, path)
        except Exception as e:
            print("Error when change vc file: ", e)
    
    path_file_match_jikken_trigger = os.path.join(link_folder_output, 'match_jikken_trigger.json')
    with open(path_file_match_jikken_trigger, 'w', encoding='utf-8') as json_file:
        json.dump(dict_match_jikken_trigger, json_file, ensure_ascii=False, indent=4)  