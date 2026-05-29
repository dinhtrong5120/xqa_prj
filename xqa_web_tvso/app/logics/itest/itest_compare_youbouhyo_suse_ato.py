import pandas as pd
import shutil
import os
from openpyxl import load_workbook
import json
from xqa_web_tvso.app.setting import *
from xqa_web_tvso.app.logics.itest.itest_btn_man import get_link_folder_output, data_normalization
from xqa_web_tvso.app.logics.itest.create_youbohyo_cpfile import *

def copy_youbouhyo(list_option_base, option_taiso):
    car_name_taiso = option_taiso.split('_')[0]
    folder_output_itest = get_link_folder_output(list_option_base, option_taiso)
    if folder_output_itest == "":
        file_source = ""
        file_json_trigger = ""
    else:
        file_source = os.path.join(folder_output_itest, "compare_配車要望表.xlsx")
        file_json_trigger = os.path.join(folder_output_itest, "match_jikken_trigger.json")

    if os.path.exists(file_source) == False or os.path.exists(file_json_trigger) == False:
        message = "Please run step 1 before running step 2"
        file_destination = ""
    else:
        message = ""
        file_name = f"{car_name_taiso}_2番手_削減後_Car配車要望表.xlsx"
        file_destination = os.path.join(folder_output_itest, file_name)
        shutil.copy(file_source, file_destination)

    return file_json_trigger,file_destination, message


def return_compare_file(file_json_trigger,file_compare_path):
    wb = load_workbook(file_compare_path)
    list_sheet = wb.sheetnames
    list_update = {}

    with open(file_json_trigger, 'r', encoding='utf-8') as file:
        data = json.load(file)

    records = []
    for key, values in data.items():
        index_ = key.rfind('_')
        name_file = key[0:index_]
        for value in values:
            records.append({'ファイル名': str(name_file).lower(), '実験_json': value})

    df_json = pd.DataFrame(records)
    df_json["flag"] = "delete"
    df_json['jikken_df_youbohyo'] = df_json['実験_json'].apply(data_normalization)
    df_json['ファイル名_'] = df_json['ファイル名'].apply(data_normalization)
    df_json['jikken_df_youbohyo'] = df_json['jikken_df_youbohyo'].str.lower()

    for name_ws in list_sheet:
        ws = wb[name_ws]

        df_youbohyo = pd.read_excel(file_compare_path, sheet_name=name_ws, header=18)
        df_youbohyo['jikken_df_youbohyo'] = df_youbohyo['実験'].apply(data_normalization)
        df_youbohyo['ファイル名_'] = df_youbohyo['ファイル名'].apply(data_normalization)
        df_youbohyo['ファイル名_'] = df_youbohyo['ファイル名_'].str.lower()
        df_youbohyo['jikken_df_youbohyo'] = df_youbohyo['jikken_df_youbohyo'].str.lower()

        merged_df = pd.merge(df_youbohyo, df_json, on=["ファイル名_", "jikken_df_youbohyo"], how='left')
        columns_to_drop = ['jikken_df_youbohyo', 'ファイル名_', 'ファイル名_y', '実験_json']
        merged_df = merged_df.drop(columns=columns_to_drop)

        df_delete = merged_df.loc[(merged_df['status'] == 1) & (merged_df['flag'] == 'delete')]
        df_keep = merged_df.loc[~((merged_df['status'] == 1) & (merged_df['flag'] == 'delete'))]
        df_delete = df_delete.drop(columns='flag')

        list_update[name_ws] = (df_keep, df_delete)

        for row in ws.iter_rows(min_row=20, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.value = None

        df_delete_start_row = 19 + len(df_keep)
        ws.cell(row=df_delete_start_row + 4, column=1, value="削減可↓")
        wb.save(file_compare_path)

    with pd.ExcelWriter(file_compare_path, engine='openpyxl', mode='a', if_sheet_exists="overlay") as writer:
        for key, (df0, df1) in list_update.items():
            df0.to_excel(writer, sheet_name=key, startrow=19, index=False, header=False)
            df_delete_start_row = 19 + len(df0)
            df1.to_excel(writer, sheet_name=key, startrow=df_delete_start_row + 5, index=False, header=False)


def run_compare_youbouhyo_suse_ato(list_option_base, option_taiso):
    file_json_trigger,file_compare_path, message = copy_youbouhyo(list_option_base, option_taiso)
    if message == "":
        # Create youbohyo compare file 
        file_name = "Car配車要望比較.xlsx"
        output_path = get_link_folder_output(list_option_base, option_taiso)
        youbohyo_compare_output = os.path.join(output_path, file_name)
        cp_youbo_file(list_option_base, option_taiso, file_json_trigger, youbohyo_compare_output)
        # Create 2rd delete jikken file 
        return_compare_file(file_json_trigger,file_compare_path)
        message = "Completed!!!"
    else:
        pass
    return message
